"""Bounded, local document parser. stdout is a JSON-lines progress protocol."""

# 学习导读：Node 的 document-parser.ts 启动本程序，读取 JSON 进度并在结束后加载 result.json。
# 解析器只产生正文、位置和质量警告；是否发布为可检索版本，由 Node 业务层决定。
import argparse
import contextlib
import io
import json
import math
import os
from pathlib import Path
import re
import subprocess
import sys
import zipfile

import pdfplumber
import pypdfium2 as pdfium
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from PIL import Image, ImageStat

try:
    import numpy as np
    from rapidocr_onnxruntime import RapidOCR
except ImportError:
    np = None
    RapidOCR = None


MAX_BYTES = 20 * 1024 * 1024
MAX_TEXT = 1024 * 1024
MAX_BLOCKS = 12000
Image.MAX_IMAGE_PIXELS = 40_000_000


class ParseFailure(Exception):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code


def emit(event):
    print(json.dumps(event, ensure_ascii=False), flush=True)


def clean(text):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text).replace("\r\n", "\n").replace("\r", "\n").strip()


def normalized_box(box, width, height):
    return [max(0, min(1, float(value) / size)) for value, size in zip(box, [width, height, width, height])]


class Parser:
    def __init__(self, options):
        # 三类结果分别供分块、原文预览和人工复核使用；不在这里连接业务数据库。
        self.options = options
        self.blocks = []
        self.pages = []
        self.warnings = []
        self.engine = None
        self.text_bytes = 0

    def warning(self, page, code, message):
        if not any(item["page"] == page and item["code"] == code for item in self.warnings):
            self.warnings.append({"page": page, "code": code, "message": message})

    def block(self, text, page, box, method, label, kind="paragraph", heading=None, confidence=None):
        # 所有格式最终都转换成同一种正文块，便于 Node 统一分块和关联来源。
        text = clean(text)
        if not text:
            return
        self.text_bytes += len(text.encode("utf-8"))
        if self.text_bytes > MAX_TEXT or len(self.blocks) >= MAX_BLOCKS:
            raise ParseFailure("document_too_complex", "提取正文过多，请拆分文档后上传")
        identifier = f"block-{len(self.blocks) + 1}"
        self.blocks.append({
            "id": identifier, "text": text, "heading": heading, "kind": kind,
            "confidence": confidence,
            "location": {"blockId": identifier, "page": page, "box": box, "method": method, "label": label},
        })

    def ocr(self, image, page, label, image_location=None):
        # off 不代表图片已经被成功读取，需要给出警告，让后续版本进入人工复核。
        if self.options.ocr == "off":
            self.warning(page, "ocr_required", "检测到图片内容，启用本地 OCR 后重试，或核对原件并补充正文")
            return
        if RapidOCR is None:
            raise ParseFailure("ocr_unavailable", "缺少本地 OCR 依赖，请安装 parser/requirements.txt")
        if image.width * image.height > self.options.max_pixels:
            image.thumbnail((int(math.sqrt(self.options.max_pixels)),) * 2)
        if ImageStat.Stat(image.convert("L")).stddev[0] < 1:
            return
        ink = np.asarray(image.convert("L")) < 160
        if int((ink.mean(axis=1) > 0.45).sum()) > 5 and int((ink.mean(axis=0) > 0.4).sum()) > 5:
            self.warning(page, "scan_table", "扫描图片可能包含表格，请核对行列对应关系")
        # 使用随库提供的 ONNX 权重在本机推理；进度协议不能被 OCR 库的 stdout 日志打断。
        with contextlib.redirect_stdout(sys.stderr):
            if self.engine is None:
                self.engine = RapidOCR(intra_op_num_threads=2, inter_op_num_threads=1, text_score=0.0)
            result, _ = self.engine(np.asarray(image.convert("RGB")), text_score=0.0)
            rotated = False
            # 文本行分类器处理倒置；效果较差时再尝试 90 度旋转，改善横置扫描页的识别。
            score = sum(len(item[1]) * float(item[2]) for item in result or [])
            if not result or score / max(1, sum(len(item[1]) for item in result)) < 0.75:
                alternative, _ = self.engine(np.asarray(image.rotate(90, expand=True).convert("RGB")), text_score=0.0)
                if sum(len(item[1]) * float(item[2]) for item in alternative or []) > score:
                    result, rotated = alternative, True
        if not result:
            self.warning(page, "unreadable_image", "页面含图片，但未识别到可用文字，请核对原件")
            return
        if len(result) > 2500:
            raise ParseFailure("document_too_complex", "单页文字块过多，请拆分文档")
        low_quality = False
        for points, text, confidence in result:
            confidence = float(confidence)
            low_quality |= confidence < 0.85
            if rotated:
                # 检测框要映射回原图坐标，前端才可在未旋转的 PDF 页面上高亮来源。
                points = [[image.width - point[1], point[0]] for point in points]
            box = normalized_box(
                [min(p[0] for p in points), min(p[1] for p in points),
                 max(p[0] for p in points), max(p[1] for p in points)],
                image.width, image.height,
            )
            self.block(text, page, image_location or (box if page else None), "ocr", label, confidence=confidence)
        # 阈值只是复核提示，不代表数学意义上的文本正确率。
        if low_quality:
            self.warning(page, "low_ocr_confidence", "部分文字识别置信度较低，请核对数字、名称和表格")
        if rotated:
            self.warning(page, "rotated_scan", "已识别横置扫描页，请确认阅读顺序")
        if page is None:
            self.warning(None, "embedded_image", "Word 内嵌图片已识别，请核对其文字与正文的关系")

    def pdf(self, path):
        # 文字提取与光栅渲染分别由专门库负责；按页决定哪条路径适合当前内容。
        try:
            document = pdfplumber.open(path, password="")
            raster = pdfium.PdfDocument(path)
        except Exception as error:
            raise ParseFailure("invalid_pdf", "PDF 无法读取，可能已加密或损坏；请上传未加密的有效文件") from error
        try:
            if not 1 <= len(raster) <= self.options.max_pages:
                raise ParseFailure("too_many_pages", f"PDF 页数须为 1 到 {self.options.max_pages} 页")
            for index, page in enumerate(document.pages):
                number = index + 1
                emit({"type": "progress", "stage": "extracting", "processedPages": index, "totalPages": len(document.pages)})
                width, height = float(page.width), float(page.height)
                if width <= 0 or height <= 0 or width * height > 20_000_000:
                    raise ParseFailure("page_too_large", "PDF 页面尺寸过大，请使用正常纸张尺寸重新导出")
                text = page.extract_text() or ""
                image_area = sum(max(0, item["x1"] - item["x0"]) * max(0, item["bottom"] - item["top"]) for item in page.images)
                bad_text = "\ufffd" in text or "(cid:" in text
                # 有页码文字的扫描页仍可能需要 OCR，不能仅凭“存在文字层”跳过图片正文。
                needs_ocr = self.options.ocr == "always" or bad_text or image_area > width * height * 0.12
                if not text.strip() and (page.images or page.curves or page.rects or page.lines):
                    needs_ocr = True
                method = "ocr" if needs_ocr else ("text" if text.strip() else "blank")
                self.pages.append({"number": number, "width": width, "height": height, "method": method})
                if needs_ocr:
                    emit({"type": "progress", "stage": "ocr", "processedPages": index, "totalPages": len(document.pages)})
                    for item in page.images:
                        dimensions = item.get("srcsize") or (0, 0)
                        if dimensions[0] * dimensions[1] > 80_000_000:
                            raise ParseFailure("image_too_large", "PDF 内嵌图片像素过多，请压缩扫描件")
                    rendered = raster[index]
                    bitmap = rendered.render(scale=min(300 / 72, math.sqrt(self.options.max_pixels / (width * height))))
                    try:
                        image = bitmap.to_pil().convert("RGB")
                        self.ocr(image, number, f"第 {number} 页")
                    finally:
                        bitmap.close()
                        rendered.close()
                    if len(page.lines) + len(page.rects) > 12:
                        self.warning(number, "scan_table", "扫描页可能包含表格，请核对行列对应关系")
                elif text.strip():
                    # 先抽表格，再排除落在表格区域内的普通文字行，避免正文重复。
                    tables = page.find_tables()
                    areas = [table.bbox for table in tables]
                    items = []
                    for table in tables:
                        rows = table.extract()
                        table_text = "\n".join(" | ".join(clean(cell or "").replace("\n", " / ") for cell in row) for row in rows)
                        items.append((table.bbox[1], table.bbox[0], table_text, "table", table.bbox))
                    for line in page.extract_text_lines(return_chars=False):
                        cx, cy = (line["x0"] + line["x1"]) / 2, (line["top"] + line["bottom"]) / 2
                        if any(x0 <= cx <= x1 and top <= cy <= bottom for x0, top, x1, bottom in areas):
                            continue
                        items.append((line["top"], line["x0"], line["text"], "paragraph",
                                      (line["x0"], line["top"], line["x1"], line["bottom"])))
                    for _, _, body, kind, box in sorted(items, key=lambda item: (item[0], item[1])):
                        self.block(body, number, normalized_box(box, width, height), "text", f"第 {number} 页", kind)
                emit({"type": "progress", "stage": "extracting", "processedPages": number, "totalPages": len(document.pages)})
                page.close()
        finally:
            document.close()
            raster.close()

    def docx(self, path):
        # DOCX 是 ZIP 容器；交给 OOXML 库前先检查总量与 XML，不把 ZIP 中的路径解压到磁盘。
        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(item.file_size for item in entries) > 64 * 1024 * 1024:
                    raise ParseFailure("archive_too_large", "Word 解压后过大，请拆分文档")
                for item in entries:
                    if item.flag_bits & 1:
                        raise ParseFailure("encrypted_document", "请上传未加密的 Word 文档")
                    if item.filename.endswith(".xml"):
                        data = archive.read(item)
                        if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
                            raise ParseFailure("invalid_docx", "不支持包含自定义实体的 Word 文档")
                        if b"txbxContent" in data or b"<w:ins " in data or b"<w:del " in data:
                            self.warning(None, "complex_word_content", "Word 包含文本框或修订记录，请核对完整内容，必要时先接受修订或导出 PDF")
                if "word/document.xml" not in archive.namelist():
                    raise ParseFailure("invalid_docx", "文件不是有效的 DOCX 文档")
            document = Document(path)
        except ParseFailure:
            raise
        except Exception as error:
            raise ParseFailure("invalid_docx", "Word 文件无法读取，请检查格式或重新导出 DOCX") from error
        heading = None
        seen_images = set()
        # DOCX 的显示分页依赖排版环境，用正文块标签定位，不编造稳定页码。
        for index, item in enumerate(document.iter_inner_content()):
            if index > MAX_BLOCKS:
                raise ParseFailure("document_too_complex", "Word 段落过多，请拆分文档")
            if index % 30 == 0:
                emit({"type": "progress", "stage": "extracting", "processedPages": 0, "totalPages": 0})
            label = f"第 {index + 1} 个正文块"
            if isinstance(item, Table):
                body = "\n".join(" | ".join(clean(cell.text).replace("\n", " / ") for cell in row.cells) for row in item.rows)
                self.block(body, None, None, "docx", label, "table", heading)
                xml = item._tbl
            else:
                is_heading = bool(item.style and (item.style.name.startswith("Heading") or item.style.name == "Title"))
                if is_heading:
                    heading = item.text.strip() or heading
                self.block(item.text, None, None, "docx", label, "heading" if is_heading else "paragraph", heading)
                xml = item._p
            for picture in xml.xpath(".//a:blip"):
                relationship = picture.get(qn("r:embed"))
                if not relationship or relationship in seen_images:
                    continue
                seen_images.add(relationship)
                part = document.part.related_parts.get(relationship)
                if not part:
                    continue
                try:
                    with Image.open(io.BytesIO(part.blob)) as embedded:
                        if embedded.width * embedded.height > 40_000_000:
                            raise ParseFailure("image_too_large", "Word 图片像素过多")
                        self.ocr(embedded.convert("RGB"), None, f"{label}中的图片")
                except ParseFailure:
                    raise
                except Exception:
                    self.warning(None, "unsupported_image", "部分 Word 图片无法识别，请核对原件")
        for section in document.sections:
            for name, part in [("页眉", section.header), ("页脚", section.footer)]:
                for paragraph in part.paragraphs:
                    text = paragraph.text.strip()
                    if text and not any(block["text"] == text for block in self.blocks):
                        self.block(text, None, None, "docx", name)
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            if any(name in names for name in ["word/footnotes.xml", "word/endnotes.xml"]) or any("embeddings/" in name for name in names):
                self.warning(None, "additional_content", "文档包含脚注、尾注或嵌入对象，请核对正文是否完整")

    def legacy_word(self, path, output):
        # 老 DOC 先在独立配置目录内转换成 PDF，再复用 PDF 路径；转换后仍要求人工核对。
        executable = self.options.libreoffice
        if not executable or not Path(executable).is_file():
            raise ParseFailure("converter_unavailable", "旧版 DOC 需要配置 KNOWLEDGE_LIBREOFFICE_PATH，或先另存为 DOCX")
        profile = output / "office-profile"
        profile.mkdir()
        try:
            subprocess.run([
                executable, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--norestore",
                "--convert-to", "pdf", "--outdir", str(output), str(path),
            ], cwd=output, stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL, timeout=60, check=True,
                creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
            converted = output / f"{path.stem}.pdf"
            if not converted.is_file() or converted.stat().st_size > MAX_BYTES:
                raise ValueError("Invalid conversion result")
            converted.replace(output / "preview.pdf")
        except Exception as error:
            raise ParseFailure("conversion_failed", "DOC 转换失败，请另存为 DOCX 或 PDF 后上传") from error
        self.pdf(output / "preview.pdf")
        self.warning(None, "converted_word", "旧版 Word 已生成固定 PDF 预览，请确认转换后的内容与分页")


def main():
    # 命令行参数来自 Node。输出只有进度/终态与受限的 JSON 文件，原件不被改写。
    arguments = argparse.ArgumentParser()
    arguments.add_argument("--input", required=True)
    arguments.add_argument("--output", required=True)
    arguments.add_argument("--type", required=True)
    arguments.add_argument("--max-pages", type=int, default=120)
    arguments.add_argument("--max-pixels", type=int, default=12_000_000)
    arguments.add_argument("--ocr", choices=["auto", "always", "off"], default="auto")
    arguments.add_argument("--libreoffice", default="")
    options = arguments.parse_args()
    path, output = Path(options.input).resolve(), Path(options.output).resolve()
    if not 1 <= path.stat().st_size <= MAX_BYTES:
        raise ParseFailure("file_too_large", "文件为空或超过 20 MiB")
    if not 1 <= options.max_pages <= 500 or not 1_000_000 <= options.max_pixels <= 20_000_000:
        raise ParseFailure("invalid_limits", "解析资源限制无效")
    parser = Parser(options)
    if options.type == "application/pdf":
        parser.pdf(path)
    elif options.type.endswith("wordprocessingml.document"):
        parser.docx(path)
    elif options.type == "application/msword":
        parser.legacy_word(path, output)
    else:
        raise ParseFailure("unsupported_file_type", "解析器不支持该文件类型")
    if not parser.blocks and not parser.warnings:
        raise ParseFailure("empty_document", "文档没有可处理的正文")
    (output / "result.json").write_text(json.dumps({
        "pages": parser.pages, "blocks": parser.blocks, "warnings": parser.warnings,
    }, ensure_ascii=False), encoding="utf-8")
    emit({"type": "done"})


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    try:
        main()
    except ParseFailure as failure:
        emit({"type": "error", "code": failure.code, "message": str(failure)})
        sys.exit(1)
    except Exception:
        # Parser libraries may include document text or local paths in exceptions.
        emit({"type": "error", "code": "parse_failed", "message": "文档解析失败，请检查文件内容或重试"})
        sys.exit(1)
