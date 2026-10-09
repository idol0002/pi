"""Generate small, synthetic parser fixtures. Never uses company documents."""
import argparse
import io
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pdfencrypt import StandardEncryption
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


def scan_image():
    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)
    chinese = Path("C:/Windows/Fonts/msyh.ttc")
    font = ImageFont.truetype(str(chinese if chinese.exists() else "DejaVuSans.ttf"), 44)
    lines = [
        "Knowledge upload acceptance",
        "Release window: Friday 14:30",
        "Owner: Engineering Team",
        "青禾项目验收时间为每周五14:30" if chinese.exists() else "Acceptance: Friday 14:30",
    ]
    for index, line in enumerate(lines):
        draw.text((90, 180 + index * 100), line, font=font, fill="black")
    return image


def make_pdf(path, mode, image):
    document = canvas.Canvas(str(path), pagesize=(612, 792), encrypt=StandardEncryption("locked-fixture") if mode == "encrypted" else None)
    document.setTitle("Synthetic knowledge parser fixture")
    if mode in ("text", "mixed", "encrypted"):
        document.setFont("Helvetica", 17)
        document.drawString(48, 720, "Deployment handbook")
        document.setFont("Helvetica", 12)
        document.drawString(48, 680, "The release owner is the Engineering Team.")
        document.drawString(48, 650, "Textpage verification marker: atlasnative.")
        if mode == "mixed":
            document.showPage()
    if mode in ("scan", "mixed", "rotated"):
        source = image.rotate(90, expand=True) if mode == "rotated" else image
        document.drawImage(ImageReader(source), 35, 60, width=542, height=690, preserveAspectRatio=True, anchor="c")
        document.setFont("Helvetica", 9)
        document.drawString(48, 30, "Page 2" if mode == "mixed" else "Page 1")
    document.showPage()
    document.save()


def main():
    args = argparse.ArgumentParser()
    args.add_argument("--kind", choices=["pdf", "docx"], required=True)
    args.add_argument("--output", required=True)
    options = args.parse_args()
    output = Path(options.output)
    output.mkdir(parents=True, exist_ok=True)
    image = scan_image()
    if options.kind == "pdf":
        for mode in ["text", "scan", "mixed", "rotated", "blank", "encrypted"]:
            make_pdf(output / (mode + ".pdf"), mode, image)
    else:
        for pictures in [False, True]:
            document = Document()
            document.add_heading("Deployment Handbook", 0)
            document.add_paragraph("This synthetic document verifies upload, table extraction and citations.")
            document.add_heading("Release window", level=1)
            document.add_paragraph("The deployment window is Friday 14:30. Marker: atlasdocx.")
            table = document.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "Team"
            table.rows[0].cells[1].text = "Responsibility"
            row = table.add_row()
            row.cells[0].text = "Engineering"
            row.cells[1].text = "Release approval"
            if pictures:
                buffer = io.BytesIO()
                image.save(buffer, format="PNG")
                buffer.seek(0)
                document.add_picture(buffer, width=Inches(4))
            document.save(output / ("word-images.docx" if pictures else "word.docx"))


if __name__ == "__main__":
    main()
